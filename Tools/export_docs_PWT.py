#!/usr/bin/env python3
"""export_docs_PWT.py — regenerate the shareable PDF + Word from the build log.

Master file:  docs/BUILD_LOG_PWT.md   (edit this)
Outputs:      docs/BUILD_LOG_PWT.pdf
              docs/BUILD_LOG_PWT.docx

Pure-Python (no LibreOffice needed) so it runs anywhere, including your Windows
machine:
    python -m pip install reportlab python-docx

Usage (from the repo root):
    python Tools/export_docs_PWT.py
    python Tools/export_docs_PWT.py --src docs/BUILD_LOG_PWT.md

It parses the (project-controlled) Markdown into blocks — headings, paragraphs,
bullet/numbered lists, tables, fenced code, blockquotes, rules — and renders each
to both PDF and DOCX. Inline **bold** and `code` are honoured.
"""
import argparse
import re
import sys
from html import escape
from pathlib import Path

# ---------------------------------------------------------------- markdown parse


def parse_blocks(md: str):
    """Return a list of (kind, payload) blocks from a controlled Markdown doc."""
    lines = md.replace("\r\n", "\n").split("\n")
    blocks, i, n = [], 0, len(lines)
    para: list[str] = []

    def flush_para():
        if para:
            blocks.append(("p", " ".join(para).strip()))
            para.clear()

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            flush_para()
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            blocks.append(("code", "\n".join(code)))
            continue

        # heading
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_para()
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue

        # horizontal rule
        if re.match(r"^(---+|\*\*\*+)$", stripped):
            flush_para()
            blocks.append(("hr", ""))
            i += 1
            continue

        # table (header row + separator row of dashes)
        if stripped.startswith("|") and i + 1 < n and re.match(
            r"^\|?[\s:|-]+\|?$", lines[i + 1].strip()
        ) and "-" in lines[i + 1]:
            flush_para()
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            cells = [
                [c.strip() for c in r.strip("|").split("|")] for r in rows
            ]
            header = cells[0]
            body = cells[2:] if len(cells) > 2 else []  # skip separator row
            blocks.append(("table", (header, body)))
            continue

        # blockquote
        if stripped.startswith(">"):
            flush_para()
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(quote).strip()))
            continue

        # unordered list
        if re.match(r"^[-*]\s+", stripped):
            flush_para()
            items = []
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ul", items))
            continue

        # ordered list
        if re.match(r"^\d+\.\s+", stripped):
            flush_para()
            items = []
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("ol", items))
            continue

        # blank line -> paragraph break
        if not stripped:
            flush_para()
            i += 1
            continue

        para.append(stripped)
        i += 1

    flush_para()
    return blocks


# ---------------------------------------------------------------- PDF (reportlab)


def render_pdf(blocks, out_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable, ListFlowable, ListItem, Paragraph, Preformatted,
        SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    ACCENT = colors.HexColor("#c0392b")
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=10, leading=14,
                          alignment=TA_LEFT, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=19, spaceAfter=10,
                        textColor=ACCENT)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=14, spaceBefore=12,
                        spaceAfter=6, textColor=colors.HexColor("#222222"))
    h3 = ParagraphStyle("h3", parent=ss["Heading3"], fontSize=11.5, spaceBefore=8,
                        spaceAfter=4)
    code = ParagraphStyle("code", parent=ss["Code"], fontName="Courier",
                          fontSize=8, leading=10, backColor=colors.HexColor("#f4f4f4"),
                          borderColor=colors.HexColor("#dddddd"), borderWidth=0.5,
                          borderPadding=5, spaceAfter=6)
    quote = ParagraphStyle("quote", parent=body, leftIndent=12,
                           textColor=colors.HexColor("#444444"),
                           backColor=colors.HexColor("#fbeeec"), borderPadding=5,
                           spaceAfter=6)

    def inline(text: str) -> str:
        text = escape(text)
        text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
        text = re.sub(r"`(.+?)`", r'<font face="Courier">\1</font>', text)
        text = re.sub(r"\[(.+?)\]\((.+?)\)", r'<link href="\2"><u>\1</u></link>', text)
        return text

    hmap = {"h1": h1, "h2": h2, "h3": h3, "h4": h3, "h5": h3, "h6": h3}
    flow = []
    for kind, payload in blocks:
        if kind in hmap:
            flow.append(Paragraph(inline(payload), hmap[kind]))
        elif kind == "p":
            flow.append(Paragraph(inline(payload), body))
        elif kind == "code":
            flow.append(Preformatted(payload or " ", code))
        elif kind == "quote":
            flow.append(Paragraph(inline(payload), quote))
        elif kind == "hr":
            flow.append(Spacer(1, 4))
            flow.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc")))
            flow.append(Spacer(1, 4))
        elif kind in ("ul", "ol"):
            items = [ListItem(Paragraph(inline(it), body), leftIndent=14)
                     for it in payload]
            flow.append(ListFlowable(
                items, bulletType="bullet" if kind == "ul" else "1",
                start="circle" if kind == "ul" else "1"))
        elif kind == "table":
            header, rows = payload
            data = [[Paragraph(inline(c), body) for c in header]]
            data += [[Paragraph(inline(c), body) for c in r] for r in rows]
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f0f0")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bbbbbb")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            flow.append(t)
            flow.append(Spacer(1, 6))

    doc = SimpleDocTemplate(str(out_path), pagesize=LETTER,
                            topMargin=0.8 * inch, bottomMargin=0.8 * inch,
                            leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            title=out_path.stem)
    doc.build(flow)


# ---------------------------------------------------------------- DOCX (python-docx)


def render_docx(blocks, out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    def add_inline(paragraph, text: str):
        # split on **bold** and `code`, keep the delimiters
        for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                paragraph.add_run(part[2:-2]).bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = paragraph.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            else:
                paragraph.add_run(part)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for kind, payload in blocks:
        if kind == "h1":
            h = doc.add_heading(payload, level=0)
        elif kind in ("h2", "h3", "h4", "h5", "h6"):
            doc.add_heading(payload, level=min(int(kind[1]) - 1, 4))
        elif kind == "p":
            add_inline(doc.add_paragraph(), payload)
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(payload)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
        elif kind == "quote":
            add_inline(doc.add_paragraph(style="Intense Quote"), payload)
        elif kind == "hr":
            doc.add_paragraph("_" * 40)
        elif kind == "ul":
            for it in payload:
                add_inline(doc.add_paragraph(style="List Bullet"), it)
        elif kind == "ol":
            for it in payload:
                add_inline(doc.add_paragraph(style="List Number"), it)
        elif kind == "table":
            header, rows = payload
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, c in enumerate(header):
                table.rows[0].cells[j].paragraphs[0].add_run(c).bold = True
            for r in rows:
                cells = table.add_row().cells
                for j, c in enumerate(r[:len(header)]):
                    add_inline(cells[j].paragraphs[0], c)
    doc.save(str(out_path))


# ---------------------------------------------------------------- main


def main() -> None:
    repo = Path(__file__).resolve().parent.parent
    p = argparse.ArgumentParser(description="Export a Markdown doc to PDF + DOCX.")
    p.add_argument("--src", default="docs/BUILD_LOG_PWT.md",
                   help="Markdown file to export (relative to repo root).")
    args = p.parse_args()

    src = (repo / args.src).resolve()
    if not src.exists():
        sys.exit(f"ERROR: source not found: {src}")

    blocks = parse_blocks(src.read_text(encoding="utf-8"))
    pdf_path = src.with_suffix(".pdf")
    docx_path = src.with_suffix(".docx")

    render_pdf(blocks, pdf_path)
    print(f"wrote {pdf_path.relative_to(repo)}")
    render_docx(blocks, docx_path)
    print(f"wrote {docx_path.relative_to(repo)}")


if __name__ == "__main__":
    main()
